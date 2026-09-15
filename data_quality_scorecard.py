from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


# ---------------------------------------------------------
# Create document
# ---------------------------------------------------------

document = Document()

# Reduce page margins to keep the scorecard to one page
section = document.sections[0]
section.top_margin = Pt(25)
section.bottom_margin = Pt(25)
section.left_margin = Pt(30)
section.right_margin = Pt(30)


# ---------------------------------------------------------
# Title
# ---------------------------------------------------------

title = document.add_heading(
    "Data Quality Scorecard",
    level=1
)

title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = document.add_paragraph()

subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = subtitle.add_run(
    "Employee Dataset | Week 2 Day 6"
)

run.bold = True


# ---------------------------------------------------------
# Dataset summary
# ---------------------------------------------------------

p = document.add_paragraph()

p.add_run("Dataset: ").bold = True
p.add_run("Employee.csv    ")

p.add_run("Rows: ").bold = True
p.add_run("4,653    ")

p.add_run("Columns: ").bold = True
p.add_run("9    ")

p.add_run("Missing Values: ").bold = True
p.add_run("0")


# ---------------------------------------------------------
# DAMA dimensions
# ---------------------------------------------------------

document.add_heading(
    "DAMA Data Quality Assessment",
    level=2
)

table = document.add_table(
    rows=1,
    cols=3
)

table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"

header = table.rows[0].cells

header[0].text = "Dimension"
header[1].text = "Score / 5"
header[2].text = "Assessment"


scores = [
    (
        "Accuracy",
        "4/5",
        "Values follow realistic ranges; validation rules were applied."
    ),
    (
        "Completeness",
        "5/5",
        "No missing values were found across the 9 columns."
    ),
    (
        "Consistency",
        "4/5",
        "Categorical values are standardized and validated."
    ),
    (
        "Timeliness",
        "3/5",
        "Validation should be performed before each BI refresh."
    ),
    (
        "Validity",
        "4/5",
        "Range and allowed-value checks detect invalid data."
    ),
    (
        "Uniqueness",
        "2/5",
        "1,889 duplicate rows were identified and require review."
    ),
]


for dimension, score, assessment in scores:

    cells = table.add_row().cells

    cells[0].text = dimension
    cells[1].text = score
    cells[2].text = assessment

    for cell in cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


# ---------------------------------------------------------
# Overall score
# ---------------------------------------------------------

document.add_paragraph()

overall = document.add_paragraph()

overall_run = overall.add_run(
    "Overall Data Quality Score: 22/30 (73.3%)"
)

overall_run.bold = True
overall_run.font.size = Pt(12)


# ---------------------------------------------------------
# Validation results
# ---------------------------------------------------------

document.add_heading(
    "Great Expectations Validation",
    level=2
)

validation = document.add_paragraph()

validation.add_run(
    "8 expectations were configured. "
).bold = False

validation.add_run(
    "5 passed and 3 failed"
).bold = True

validation.add_run(
    " after intentionally introducing three test errors: "
    "invalid PaymentTier, invalid Age, and invalid Gender. "
    "All three intentional errors were detected successfully."
)


# ---------------------------------------------------------
# Key issues
# ---------------------------------------------------------

document.add_heading(
    "Key Data Quality Issues",
    level=2
)

issues = [
    "1,889 duplicate records were identified in the original dataset.",
    "Invalid PaymentTier, Age, and Gender values were detected during testing.",
    "The dataset requires automated validation before BI dashboard refreshes.",
]


for issue in issues:

    paragraph = document.add_paragraph(
        style="List Bullet"
    )

    paragraph.add_run(issue)


# ---------------------------------------------------------
# Recommendations
# ---------------------------------------------------------

document.add_heading(
    "Improvement Recommendations",
    level=2
)

recommendations = [
    "Remove or investigate duplicate records before BI consumption.",
    "Run Great Expectations validation automatically before each dashboard refresh.",
    "Enforce valid ranges and allowed categorical values at the data ingestion stage.",
    "Use aggregated or masked employee attributes when detailed personal information is unnecessary.",
]


for recommendation in recommendations:

    paragraph = document.add_paragraph(
        style="List Bullet"
    )

    paragraph.add_run(recommendation)


# ---------------------------------------------------------
# Footer
# ---------------------------------------------------------

footer = section.footer.paragraphs[0]

footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

footer.add_run(
    "Producer: Aishwarya | Consumer: BI Dashboard | "
    "Validation Framework: Great Expectations"
)


# ---------------------------------------------------------
# Save document
# ---------------------------------------------------------

output_file = "data_quality_scorecard.docx"

document.save(output_file)

print(
    f"Scorecard created successfully: {output_file}"
)