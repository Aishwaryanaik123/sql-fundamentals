"""
Groq EDA Executive Report
Upload a CSV -> run EDA -> generate an executive summary -> download a PowerPoint.
Works with Streamlit. If GROQ_API_KEY is absent, a deterministic local summary is generated.
"""
from pathlib import Path
import os
import io
import pandas as pd
import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

st.set_page_config(page_title="Groq EDA Executive Report", page_icon="📊", layout="wide")
st.title("📊 Groq EDA Executive Report")
st.caption("Upload a CSV, inspect automated EDA insights, generate an executive summary, and export a PowerPoint.")

uploaded = st.file_uploader("Upload a CSV file", type=["csv"])
if uploaded is None:
    st.info("Upload your CSV to begin, or try the included sample_data/sales_data.csv file.")
    st.stop()

try:
    df = pd.read_csv(uploaded)
except Exception as exc:
    st.error(f"Could not read CSV: {exc}")
    st.stop()

if df.empty:
    st.error("The uploaded CSV has no rows.")
    st.stop()

st.subheader("1. Dataset overview")
c1, c2, c3 = st.columns(3)
c1.metric("Rows", f"{len(df):,}")
c2.metric("Columns", f"{len(df.columns):,}")
c3.metric("Missing cells", f"{int(df.isna().sum().sum()):,}")
st.dataframe(df.head(20), use_container_width=True)

st.subheader("2. Automated EDA")
numeric = df.select_dtypes(include="number")
insights = []
insights.append(f"The dataset contains {len(df):,} rows and {len(df.columns)} columns.")
missing = df.isna().sum()
missing_cols = missing[missing > 0].sort_values(ascending=False)
if len(missing_cols):
    col = missing_cols.index[0]
    insights.append(f"Missing values are present; {col} has the most missing cells ({int(missing_cols.iloc[0])}).")
else:
    insights.append("No missing values were detected in the uploaded dataset.")
duplicates = int(df.duplicated().sum())
insights.append(f"The dataset contains {duplicates:,} duplicate rows.")
if not numeric.empty:
    col = numeric.var(numeric_only=True).sort_values(ascending=False).index[0]
    insights.append(f"{col} has the largest numeric variance; its average is {numeric[col].mean():,.2f}.")
    col2 = numeric.mean(numeric_only=True).sort_values(ascending=False).index[0]
    insights.append(f"{col2} has the highest numeric column average at {numeric[col2].mean():,.2f}.")
else:
    insights.append("No numeric columns were detected, so numeric summaries are limited.")
for i, insight in enumerate(insights[:5], 1):
    st.write(f"**Insight {i}.** {insight}")

if not numeric.empty:
    st.write("**Numeric summary**")
    st.dataframe(numeric.describe().T, use_container_width=True)

st.subheader("3. AI executive summary")
api_key = os.getenv("GROQ_API_KEY", "").strip()
if st.button("Generate executive summary", type="primary"):
    summary = None
    if api_key:
        try:
            from groq import Groq
            prompt = (
                "Write an executive summary of approximately 200 words for a business stakeholder. "
                "Use only the supplied findings; do not invent causes, metrics, or recommendations. "
                "Explain the main patterns, data-quality caveats, business implications, and practical next steps. "
                "Clearly distinguish observations from hypotheses.\n\n"
                "Five data insights:\n" + "\n".join(f"{i+1}. {x}" for i, x in enumerate(insights[:5]))
            )
            response = Groq(api_key=api_key).chat.completions.create(
                model="llama-3.3-70b-versatile",
                messages=[{"role": "user", "content": prompt}],
                temperature=0.2,
                max_tokens=350,
            )
            summary = response.choices[0].message.content.strip()
        except Exception as exc:
            st.warning(f"Groq request failed; using the built-in fallback summary. Details: {exc}")
    if not summary:
        summary = (
            f"This dataset contains {len(df):,} records across {len(df.columns)} columns. "
            + " ".join(insights[:5])
            + " These findings provide an initial view of dataset structure and quality, but they should be interpreted "
            "in the context of how the data was collected and what each field represents. Before making business decisions, "
            "validate key fields, confirm that records cover the intended time period and population, and investigate any "
            "missing or duplicate records. Numeric averages and variance are useful screening indicators, but they do not by "
            "themselves explain causes or prove that one factor drives another. Next, compare important metrics across relevant "
            "groups or time periods, check unusual values with subject-matter owners, and connect the results to agreed business "
            "targets. Use the resulting analysis to prioritize a small number of measurable actions, assign owners, and review "
            "outcomes after implementation. Where the dataset lacks context or contains quality issues, resolve those limitations "
            "before relying on the results for high-impact decisions."
        )
    st.session_state["summary"] = summary

summary = st.session_state.get("summary")
if summary:
    st.write(summary)
    st.caption(f"{len(summary.split())} words")
    st.subheader("4. Download PowerPoint")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "EDA Executive Report"
    slide.placeholders[1].text = "Automated exploratory data analysis and business recommendations"
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Dataset Overview"
    tf = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11), Inches(4.5)).text_frame
    tf.text = f"Rows: {len(df):,}\nColumns: {len(df.columns)}\nMissing cells: {int(df.isna().sum().sum()):,}\nDuplicate rows: {duplicates:,}"
    for p in tf.paragraphs:
        p.font.size = Pt(22)
    for idx, insight in enumerate(insights[:5], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"Insight {idx}"
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11), Inches(3.8))
        box.text_frame.text = insight
        box.text_frame.paragraphs[0].font.size = Pt(24)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Recommendations"
    tf = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11), Inches(4.5)).text_frame
    tf.text = (
        "1. Validate missing values, duplicates, and unusual records.\n"
        "2. Segment key metrics by relevant business dimensions.\n"
        "3. Compare results with targets and historical periods.\n"
        "4. Assign owners and track measurable outcomes.\n"
        "5. Revisit conclusions after data quality checks."
    )
    for p in tf.paragraphs:
        p.font.size = Pt(20)
    output = io.BytesIO()
    prs.save(output)
    st.download_button("Download PowerPoint (.pptx)", data=output.getvalue(),
                       file_name="EDA_Executive_Report.pptx",
                       mime="application/vnd.openxmlformats-officedocument.presentationml.presentation")
