"""Groq-powered EDA report generator.
Upload a CSV, compute basic EDA, create five grounded insights, write an executive
summary, and export a 7-slide PowerPoint. Groq is optional: without an API key,
a local template-based summary is used.
"""
from pathlib import Path
import os
import pandas as pd
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import streamlit as st

st.set_page_config(page_title="Groq EDA Executive Report", page_icon="📊", layout="wide")
st.title("📊 Groq EDA Executive Report")
st.caption("CSV → EDA → 5 evidence-based insights → executive summary → PowerPoint")

uploaded = st.file_uploader("Upload a CSV file", type=["csv"])
if uploaded is None:
    st.info("Upload a CSV to begin, or use the sample file included in the data folder.")
    sample = Path(__file__).parent / "data" / "sample_sales.csv"
    if sample.exists() and st.button("Load included sample dataset"):
        st.session_state["sample_mode"] = True
    if not st.session_state.get("sample_mode"):
        st.stop()
    uploaded = sample

try:
    df = pd.read_csv(uploaded)
except Exception as e:
    st.error(f"Could not read the CSV: {e}")
    st.stop()

if df.empty:
    st.error("The CSV has no data rows.")
    st.stop()

st.subheader("1. Dataset overview")
c1, c2, c3, c4 = st.columns(4)
c1.metric("Rows", f"{len(df):,}")
c2.metric("Columns", f"{len(df.columns):,}")
c3.metric("Missing cells", f"{int(df.isna().sum().sum()):,}")
c4.metric("Duplicate rows", f"{int(df.duplicated().sum()):,}")
st.dataframe(df.head(20), use_container_width=True)

numeric = df.select_dtypes(include="number")
categorical = df.select_dtypes(exclude="number")
st.subheader("2. Exploratory data analysis")
if not numeric.empty:
    st.write("**Numeric summary**")
    st.dataframe(numeric.describe().T, use_container_width=True)
if not categorical.empty:
    st.write("**Categorical summary**")
    st.dataframe(categorical.describe().T, use_container_width=True)

insights = []
insights.append(f"The dataset contains {len(df):,} rows across {len(df.columns)} columns.")
missing_count = int(df.isna().sum().sum())
missing_pct = missing_count / (len(df) * len(df.columns)) * 100
insights.append(f"Data completeness: {100-missing_pct:.1f}% of cells are populated ({missing_count:,} missing cells).")
if not numeric.empty:
    means = numeric.mean(numeric_only=True).sort_values(ascending=False)
    top_col = means.index[0]
    low_col = means.index[-1]
    insights.append(f"Among numeric columns, {top_col} has the highest average ({means[top_col]:,.2f}).")
    insights.append(f"Among numeric columns, {low_col} has the lowest average ({means[low_col]:,.2f}).")
    spread = numeric.std(numeric_only=True).sort_values(ascending=False)
    insights.append(f"{spread.index[0]} shows the largest standard deviation ({spread.iloc[0]:,.2f}), indicating the greatest absolute spread.")
else:
    insights.append("No numeric columns were detected, so numeric averages and distributions are unavailable.")
if not categorical.empty:
    col = categorical.columns[0]
    counts = categorical[col].astype(str).value_counts()
    insights.append(f"In '{col}', the most frequent category is '{counts.index[0]}' ({counts.iloc[0]} rows).")
else:
    insights.append("No categorical columns were detected; segment-level analysis may require a category field.")
insights = insights[:5]
while len(insights) < 5:
    insights.append("Review data definitions and validate key fields before drawing business conclusions.")

def make_summary(insight_list):
    prompt = (
        "Write an executive summary of about 200 words for business stakeholders using ONLY "
        "the verified findings below. Explain what the findings suggest, note limitations, and "
        "give practical next steps. Do not invent causes, trends, or facts.\\n\\n"
        + "\\n".join(f"- {x}" for x in insight_list)
    )
    api_key = os.getenv("GROQ_API_KEY", "").strip()
    if api_key:
        try:
            from groq import Groq
            client = Groq(api_key=api_key)
            response = client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                messages=[
                    {"role": "system", "content": "You are a careful senior data analyst. Do not invent facts."},
                    {"role": "user", "content": prompt},
                ],
                temperature=0.2,
                max_tokens=350,
            )
            return response.choices[0].message.content.strip(), "Groq AI"
        except Exception as e:
            st.warning(f"Groq was unavailable; using a local summary instead. Details: {e}")
    summary = (
        "This analysis provides an initial view of the uploaded dataset's size, completeness, "
        "numeric fields, and category distribution. The findings below highlight the strongest "
        "descriptive signals visible in the available data. They should be treated as exploratory "
        "rather than proof of business causes or future outcomes. Before making operational "
        "decisions, confirm that each column is defined consistently, check for outliers and "
        "duplicate records, and verify that missing values are handled appropriately. Where "
        "numeric fields differ in units or scale, compare standardized measures rather than raw "
        "averages or standard deviations. If category fields are available, compare key outcomes "
        "across groups and investigate material differences with domain owners. The next step is "
        "to validate these findings against the business question, add relevant time or segment "
        "context if available, and agree on a small set of measurable actions. Re-run this report "
        "after data quality issues are addressed so stakeholders can distinguish reliable signals "
        "from limitations in the source data."
    )
    return summary, "Local fallback (no Groq key)"

st.subheader("3. Five data insights")
for i, insight in enumerate(insights, 1):
    st.markdown(f"**Insight {i}.** {insight}")

summary, summary_source = make_summary(insights)
st.subheader("4. Executive summary (~200 words)")
st.caption(f"Generated with: {summary_source}")
st.write(summary)

def create_ppt(dataframe, insight_list, executive_summary, source):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    def add_text_slide(title, body, subtitle=None):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(12), Inches(0.8))
        p = title_box.text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(28); p.font.bold = True
        if subtitle:
            sb = slide.shapes.add_textbox(Inches(0.75), Inches(1.35), Inches(11.8), Inches(0.4))
            sp = sb.text_frame.paragraphs[0]; sp.text = subtitle; sp.font.size = Pt(12)
        body_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.6))
        tf = body_box.text_frame; tf.word_wrap = True
        tf.text = body
        for p in tf.paragraphs:
            p.font.size = Pt(21)
        return slide
    add_text_slide("Executive EDA Report", f"Dataset: {len(dataframe):,} rows × {len(dataframe.columns)} columns\\nGenerated by Groq EDA Executive Report", "CSV → EDA → Insights → Recommendations")
    for i, item in enumerate(insight_list, 1):
        add_text_slide(f"Insight {i} of 5", item, "Evidence-based descriptive finding")
    add_text_slide("Recommendations", "Executive summary: " + executive_summary + "\\n\\nRecommended next steps:\\n1. Validate missing values, duplicates, and outliers.\\n2. Confirm metric definitions and units.\\n3. Segment outcomes by relevant categories or time periods.\\n4. Validate findings with stakeholders before acting.\\n5. Track agreed KPIs and refresh the report regularly.")
    out = Path(__file__).parent / "output" / "Groq_EDA_Executive_Report.pptx"
    out.parent.mkdir(exist_ok=True)
    prs.save(out)
    return out

if st.button("Generate PowerPoint"):
    ppt_path = create_ppt(df, insights, summary, summary_source)
    st.success("PowerPoint created.")
    st.download_button("⬇️ Download PowerPoint", data=ppt_path.read_bytes(),
                       file_name=ppt_path.name,
                       mime="application/vnd.openxmlformats-officedocument.presentationml.presentation")
